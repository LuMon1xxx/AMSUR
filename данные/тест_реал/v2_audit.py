# -*- coding: utf-8 -*-
"""Приёмка v2 по §5 промта (метод 1-в-1 как audit_nosvc.py от 23.09):
1) v2 Учителя-лист -> по-дням сетки (фамилия x 12 слотов) в АМСУР_движок_v2_по_дням/
2) окна без служебного: ЧЕЛОВЕК vs v2 (те же SERVICE_KW)
3) парное сравнение по общим фамилиям
4) СанПиН-аудит v2: классо-день distinct-слоты vs кэпы (5-6 кл <=6, 7-11 <=7),
   учитель <=10/день; дубли (класс,предмет,день): всего пар и % рядом
Выход — v2_audit_out.txt."""

import glob
import openpyxl
import os
import sys
from collections import Counter

BASE = "данные/тест_реал"
# §6 приёмки: тот же метод на v3 — параметром (default v2 для совместимости).
# v4 (24.09): тот же метод на ФИНАЛ_v4 (только _v4 имена; v2/v3 не тронуты).
VER = sys.argv[1] if len(sys.argv) > 1 else "v2"
assert VER in ("v2", "v3", "v4"), "usage: v2_audit.py [v2|v3|v4]"
V2X = f"{BASE}/АМСУР_движок_ФИНАЛ_{VER}.xlsx"
V2D = f"{BASE}/АМСУР_движок_{VER}_по_дням"
HUM = f"{BASE}/ЧЕЛОВЕК_школа_с_фото"
DAYS = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница"]
DSHORT = {
    "Пн": "Понедельник",
    "Вт": "Вторник",
    "Ср": "Среда",
    "Чт": "Четверг",
    "Пт": "Пятница",
}
SERVICE_KW = ("цех", "к/ч", "кч", "над", "аст", "деж", "3кл")
out = []


def is_service(v):
    if v is None:
        return True
    s = str(v).strip().lower().replace(" ", "")
    if s == "":
        return True
    return any(k in s for k in SERVICE_KW)


def surname(full):
    return str(full).strip().split()[0]


# 1. конвертер v2 -> по-дням
wb = openpyxl.load_workbook(V2X, data_only=True)
ws = wb["Учителя"]
grid = {}  # day -> {surname: {slot: class}}
days_seen = set()
for r in list(ws.iter_rows(values_only=True))[1:]:
    if not r[0]:
        continue
    t, d, slot, cls = (
        surname(r[0]),
        str(r[1]).strip(),
        int(r[2]),
        str(r[3]).strip().lower(),
    )
    days_seen.add(d)
    grid.setdefault(DSHORT[d], {}).setdefault(t, {})[slot] = cls
assert days_seen == set(DSHORT), f"дни {VER}: {days_seen}"
os.makedirs(V2D, exist_ok=True)
teachers = sorted({t for dd in grid.values() for t in dd})
for d in DAYS:
    w = openpyxl.Workbook()
    s = w.active
    s.title = d[:11]
    s.cell(1, 1, "Преподователи")
    s.cell(1, 2, d)
    for c in range(1, 13):
        s.cell(2, 1 + c, c)
    for i, t in enumerate(teachers):
        s.cell(3 + i, 1, t)
        for slot, cls in grid.get(d, {}).get(t, {}).items():
            if 1 <= slot <= 12:
                s.cell(3 + i, 1 + slot, cls)
    w.save(f"{V2D}/АМСУР_{VER}_Учителя_{d}.xlsx")
out.append(f"{VER} grids: {len(teachers)} учителей x 5 дней -> {V2D}/")


def audit(folder, pattern, label):
    files = sorted(glob.glob(folder + "/" + pattern))
    gaps = lessons = tdays = dwg = 0
    per = {}
    for f in files:
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        for r in list(ws.iter_rows(values_only=True))[2:]:
            name = str(r[0]).strip() if r[0] else ""
            if not name or "препод" in name.lower():
                continue
            cells = list(r[1:13])
            idxs = [i + 1 for i, c in enumerate(cells) if not is_service(c)]
            if not idxs:
                continue
            lessons += len(idxs)
            tdays += 1
            fl, ll = min(idxs), max(idxs)
            g = sum(1 for s in range(fl, ll + 1) if is_service(cells[s - 1]))
            gaps += g
            if g > 0:
                dwg += 1
            a = per.setdefault(name, [0, 0])
            a[0] += g
            a[1] += len(idxs)
    out.append(
        f"{label}: окон={gaps} уроков={lessons} учителе-дней={tdays} "
        f"дней_с_окнами={dwg} ({dwg / max(tdays, 1) * 100:.1f}%) на100={gaps / max(lessons, 1) * 100:.1f}"
    )
    top = sorted(per.items(), key=lambda kv: -kv[1][0])[:8]
    for k, (gg, nn) in top:
        out.append(f"   {k}: gaps={gg} n={nn}")
    return gaps, lessons, tdays, per


# 2. окна без служебного
out.append("== БЕЗ служебного (тот же метод, что 140 vs 315) ==")
hg, hl, ht, hper = audit(HUM, "ЧЕЛОВЕК_Учителя_*.xlsx", "ЧЕЛОВЕК")
vg, vl, vt, vper = audit(V2D, f"АМСУР_{VER}_Учителя_*.xlsx", f"ДВИЖОК-{VER}")


# 3. парное сравнение
def norm(t):
    return t.replace("ё", "е").replace("Ё", "Е").strip().lower()


hmap = {norm(t): t for t in hper}
vmap = {norm(t): t for t in vper}
common = sorted(set(hmap) & set(vmap))
out.append(
    f"common={len(common)} ({VER}-only={len(vmap) - len(common)}, hum-only={len(hmap) - len(common)})"
)
we = wr = tie = 0
for c in common:
    e, r = vper[vmap[c]][0], hper[hmap[c]][0]
    flag = VER if e < r else ("hum" if r < e else "=")
    if e < r:
        we += 1
    elif r < e:
        wr += 1
    else:
        tie += 1
    out.append(f"  {vmap[c]} | {VER}={e} hum={r} [{flag}]")
out.append(f"paired: {VER}_better={we} hum_better={wr} tie={tie}")


# 4. СанПиН-аудит v2 по листу Учителя (класс+предмет+день есть)
def grade_of(cls):
    d = "".join(ch for ch in cls if ch.isdigit())
    return int(d[:-1] if len(d) > 2 else d) if d else 0


wb = openpyxl.load_workbook(V2X, data_only=True)
ws = wb["Учителя"]
rows = [r for r in list(ws.iter_rows(values_only=True))[1:] if r[0]]
class_day = {}
teach_day = {}
pairs = {}
for r in rows:
    t, d, slot, cls, subj = (
        surname(r[0]),
        str(r[1]).strip(),
        int(r[2]),
        str(r[3]).strip(),
        str(r[4]).strip(),
    )
    class_day.setdefault((cls, d), set()).add(slot)
    teach_day.setdefault((t, d), []).append(slot)
    pairs.setdefault((cls, subj, d), []).append(slot)
over = []
for (cls, d), slots in sorted(class_day.items()):
    cap = 6 if grade_of(cls) in (5, 6) else 7
    if len(slots) > cap:
        over.append(f"{cls} {d}: {len(slots)}>{cap}")
out.append(f"sanpin классо-дни сверх кэпа: {len(over)} (было: школа 1, движок-v1 4)")
for o in over[:15]:
    out.append(f"   {o}")
tmax = max(len(v) for v in teach_day.values())
tover = sum(1 for v in teach_day.values() if len(v) > 10)
out.append(f"учитель макс/день={tmax}, дней с перегрузом >10: {tover}")
nd, nadj = 0, 0
for k, slots in pairs.items():
    if len(slots) == 2:
        nd += 1
        if abs(slots[0] - slots[1]) == 1:
            nadj += 1
out.append(
    f"дубли (ровно 2/день): {nd}, рядом: {nadj} ({nadj / max(nd, 1) * 100:.0f}%) "
    f"[школа: 27/89%, движок-v1: 95/28%]"
)
pup = "pupilHard==0 (gate в прогоне: hard=0, pupilHard=0, 888/888)"

open(
    f"C:/Users/C39C~1/AppData/Local/Temp/opencode/{VER}_audit_out.txt",
    "w",
    encoding="utf-8",
).write("\n".join(out) + "\n")
print("ok " + VER)
